"""Convert FastMoss product-search workbooks into a products upsert migration."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from import_fastmoss_xlsx import optional_number, parse_prices


PRODUCT_ID_PATTERN = re.compile(r"/detail/(\d+)")


def read_rows(paths: list[Path]):
    by_id = {}
    duplicates = 0
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "FastMoss" not in workbook.sheetnames:
            raise ValueError(f"Aba FastMoss ausente em {path.name}")
        sheet = workbook["FastMoss"]
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            values = values[:20]
            if not any(value is not None for value in values):
                continue
            (
                nome,
                status_origem,
                loja_nome,
                pais_regiao,
                categoria,
                preco,
                comissao,
                vendas_7_dias,
                gmv_7_dias,
                vendas_totais,
                gmv_total,
                _total_criadores,
                _conversao_criadores,
                _videos_promocionais,
                _lives_promocionais,
                imagem_url,
                fastmoss_produto_url,
                tiktok_shop_url,
                fastmoss_loja_url,
                data_upload,
            ) = values
            match = PRODUCT_ID_PATTERN.search(str(fastmoss_produto_url or ""))
            if not match:
                raise ValueError(f"ID ausente em {path.name}, linha {row_number}")
            produto_externo_id = match.group(1)
            preco_texto, preco_min, preco_max = parse_prices(preco)
            product = {
                "produto_externo_id": produto_externo_id,
                "nome": str(nome),
                "categoria": str(categoria) if categoria is not None else None,
                "imagem_url": str(imagem_url) if imagem_url is not None else None,
                "loja_nome": str(loja_nome) if loja_nome is not None else None,
                "preco_texto": preco_texto,
                "preco_min_brl": preco_min,
                "preco_max_brl": preco_max,
                "pais_regiao": str(pais_regiao) if pais_regiao is not None else None,
                "comissao_percentual": optional_number(comissao),
                "unidades_vendidas_periodo": int(vendas_7_dias) if vendas_7_dias is not None else None,
                "unidades_vendidas_total": int(vendas_totais) if vendas_totais is not None else None,
                "receita_periodo_brl": optional_number(gmv_7_dias),
                "receita_total_brl": optional_number(gmv_total),
                "data_estimada_upload": str(data_upload) if data_upload is not None else None,
                "fastmoss_produto_url": str(fastmoss_produto_url),
                "tiktok_shop_url": str(tiktok_shop_url) if tiktok_shop_url is not None else None,
                "fastmoss_loja_url": str(fastmoss_loja_url) if fastmoss_loja_url is not None else None,
                "status_origem": str(status_origem) if status_origem is not None else None,
                "fonte_dados": path.name,
            }
            if produto_externo_id in by_id:
                duplicates += 1
            by_id[produto_externo_id] = product

    products = sorted(
        by_id.values(),
        key=lambda product: (
            -(product["unidades_vendidas_periodo"] or 0),
            product["produto_externo_id"],
        ),
    )
    for index, product in enumerate(products, start=1):
        product["ranking_origem"] = index
    return products, duplicates


def migration_sql(products, source_names):
    payload = json.dumps(products, ensure_ascii=False, separators=(",", ":"))
    return f"""-- FastMoss product searches supplied by the administrator on 2026-10-18.
-- Sources: {', '.join(source_names)}.
-- Period metrics represent the last 7 days. Repeated external IDs are collapsed; the later file wins.
-- Existing UUIDs and editorial fields (status, destaque and ordem_destaque) are preserved.
WITH catalog AS (
  SELECT *
  FROM jsonb_to_recordset($catalog${payload}$catalog$::jsonb) AS item(
    produto_externo_id TEXT, nome TEXT, categoria TEXT, imagem_url TEXT,
    loja_nome TEXT, preco_texto TEXT, preco_min_brl NUMERIC, preco_max_brl NUMERIC,
    pais_regiao TEXT, comissao_percentual NUMERIC, unidades_vendidas_periodo BIGINT,
    unidades_vendidas_total BIGINT, receita_periodo_brl NUMERIC,
    receita_total_brl NUMERIC, data_estimada_upload TEXT,
    fastmoss_produto_url TEXT, tiktok_shop_url TEXT, fastmoss_loja_url TEXT,
    status_origem TEXT, fonte_dados TEXT, ranking_origem INTEGER
  )
)
INSERT INTO public.products (
  produto_externo_id, nome, categoria, imagem_url, loja_nome, preco_texto,
  preco_min_brl, preco_max_brl, pais_regiao, comissao_percentual,
  unidades_vendidas_periodo, unidades_vendidas_total,
  receita_periodo_brl, receita_total_brl, data_estimada_upload,
  fastmoss_produto_url, tiktok_shop_url, fastmoss_loja_url,
  status_origem, fonte_dados, ranking_origem, status, destaque, ordem_destaque
)
SELECT
  produto_externo_id, nome, categoria, imagem_url, loja_nome, preco_texto,
  preco_min_brl, preco_max_brl, pais_regiao, comissao_percentual,
  unidades_vendidas_periodo, unidades_vendidas_total,
  receita_periodo_brl, receita_total_brl,
  CASE WHEN data_estimada_upload IS NULL THEN NULL
    ELSE data_estimada_upload::timestamp AT TIME ZONE 'America/Sao_Paulo' END,
  fastmoss_produto_url, tiktok_shop_url, fastmoss_loja_url,
  status_origem, fonte_dados, ranking_origem,
  CASE WHEN status_origem = 'Disponível' THEN 'ativo' ELSE 'pausado' END,
  false, 0
FROM catalog
ON CONFLICT (produto_externo_id) DO UPDATE SET
  nome = EXCLUDED.nome,
  categoria = EXCLUDED.categoria,
  imagem_url = EXCLUDED.imagem_url,
  loja_nome = EXCLUDED.loja_nome,
  preco_texto = EXCLUDED.preco_texto,
  preco_min_brl = EXCLUDED.preco_min_brl,
  preco_max_brl = EXCLUDED.preco_max_brl,
  pais_regiao = EXCLUDED.pais_regiao,
  comissao_percentual = EXCLUDED.comissao_percentual,
  unidades_vendidas_periodo = EXCLUDED.unidades_vendidas_periodo,
  unidades_vendidas_total = EXCLUDED.unidades_vendidas_total,
  receita_periodo_brl = EXCLUDED.receita_periodo_brl,
  receita_total_brl = EXCLUDED.receita_total_brl,
  data_estimada_upload = EXCLUDED.data_estimada_upload,
  fastmoss_produto_url = EXCLUDED.fastmoss_produto_url,
  tiktok_shop_url = EXCLUDED.tiktok_shop_url,
  fastmoss_loja_url = EXCLUDED.fastmoss_loja_url,
  status_origem = EXCLUDED.status_origem,
  fonte_dados = EXCLUDED.fonte_dados,
  ranking_origem = EXCLUDED.ranking_origem;
"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--migration", required=True, type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    args = parser.parse_args()

    products, duplicates = read_rows(args.inputs)
    args.migration.parent.mkdir(parents=True, exist_ok=True)
    args.migration.write_text(
        migration_sql(products, [path.name for path in args.inputs]), encoding="utf-8"
    )
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(
            json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps({"products": len(products), "duplicates_collapsed": duplicates}))


if __name__ == "__main__":
    main()
