"""Convert one or more FastMoss ranking workbooks into a products upsert migration.

Usage:
  python tools/import_fastmoss_xlsx.py --migration OUTPUT.sql --json OUTPUT.json INPUT.xlsx [...]
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


COLUMNS = [
    "produto_externo_id",
    "nome",
    "categoria",
    "imagem_url",
    "loja_nome",
    "preco_texto",
    "preco_min_brl",
    "preco_max_brl",
    "pais_regiao",
    "comissao_percentual",
    "unidades_vendidas_periodo",
    "crescimento_vendas_percentual",
    "unidades_vendidas_total",
    "receita_periodo_brl",
    "receita_total_brl",
    "data_estimada_upload",
    "unidades_vendidas_total_loja",
    "fastmoss_produto_url",
    "tiktok_shop_url",
    "fastmoss_loja_url",
    "status_origem",
    "fonte_dados",
    "ranking_origem",
]


def optional_number(value):
    if value is None or value == "-":
        return None
    if isinstance(value, (int, float)):
        return value
    normalized = str(value).strip().replace("%", "").replace("R$", "").strip()
    if "," in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    return float(normalized)


def parse_prices(value):
    text = str(value).strip()
    numbers = re.findall(r"\d[\d.]*,\d{2}|\d+(?:\.\d{2})?", text)
    parsed = [optional_number(number) for number in numbers]
    if not parsed:
        return text, None, None
    return text, min(parsed), max(parsed)


def read_rows(paths: list[Path]):
    by_id = {}
    duplicates = 0
    for path in paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "FastMoss" not in workbook.sheetnames:
            raise ValueError(f"Aba FastMoss ausente em {path.name}")
        sheet = workbook["FastMoss"]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            values = values[:19]
            if not any(value is not None for value in values):
                continue
            (
                nome,
                produto_externo_id,
                imagem_url,
                loja_nome,
                preco,
                pais_regiao,
                categoria,
                comissao,
                unidades_periodo,
                crescimento,
                unidades_total,
                receita_total,
                receita_periodo,
                data_upload,
                unidades_loja,
                fastmoss_produto_url,
                tiktok_shop_url,
                fastmoss_loja_url,
                status_origem,
            ) = values
            if not produto_externo_id or not nome:
                raise ValueError(f"Linha sem ID ou nome em {path.name}")
            preco_texto, preco_min, preco_max = parse_prices(preco)
            product = {
                "produto_externo_id": str(produto_externo_id),
                "nome": str(nome),
                "categoria": str(categoria) if categoria is not None else None,
                "imagem_url": str(imagem_url) if imagem_url is not None else None,
                "loja_nome": str(loja_nome) if loja_nome is not None else None,
                "preco_texto": preco_texto,
                "preco_min_brl": preco_min,
                "preco_max_brl": preco_max,
                "pais_regiao": str(pais_regiao) if pais_regiao is not None else None,
                "comissao_percentual": optional_number(comissao),
                "unidades_vendidas_periodo": int(unidades_periodo) if unidades_periodo is not None else None,
                "crescimento_vendas_percentual": optional_number(crescimento),
                "unidades_vendidas_total": int(unidades_total) if unidades_total is not None else None,
                "receita_periodo_brl": optional_number(receita_periodo),
                "receita_total_brl": optional_number(receita_total),
                "data_estimada_upload": str(data_upload) if data_upload is not None else None,
                "unidades_vendidas_total_loja": int(unidades_loja) if unidades_loja is not None else None,
                "fastmoss_produto_url": str(fastmoss_produto_url) if fastmoss_produto_url is not None else None,
                "tiktok_shop_url": str(tiktok_shop_url) if tiktok_shop_url is not None else None,
                "fastmoss_loja_url": str(fastmoss_loja_url) if fastmoss_loja_url is not None else None,
                "status_origem": str(status_origem) if status_origem is not None else None,
                "fonte_dados": path.name,
            }
            product_id = product["produto_externo_id"]
            if product_id in by_id:
                duplicates += 1
            by_id[product_id] = product

    products = sorted(
        by_id.values(),
        key=lambda product: (
            -(product["unidades_vendidas_periodo"] or 0),
            product["produto_externo_id"],
        ),
    )
    for index, product in enumerate(products, start=1):
        product["ranking_origem"] = index
    return products, duplicates


def migration_sql(products, source_names):
    payload = json.dumps(products, ensure_ascii=False, separators=(",", ":"))
    return f"""-- FastMoss/TikTok Shop rankings supplied by the administrator on 2026-10-18.
-- Sources: {', '.join(source_names)}.
-- Duplicate external IDs across the source files are collapsed; the later file wins.
-- Re-running is safe: external product IDs are upserted and existing internal UUIDs/editorial fields are preserved.
WITH catalog AS (
  SELECT *
  FROM jsonb_to_recordset($catalog${payload}$catalog$::jsonb) AS item(
    produto_externo_id TEXT, nome TEXT, categoria TEXT, imagem_url TEXT,
    loja_nome TEXT, preco_texto TEXT, preco_min_brl NUMERIC, preco_max_brl NUMERIC,
    pais_regiao TEXT, comissao_percentual NUMERIC, unidades_vendidas_periodo BIGINT,
    crescimento_vendas_percentual NUMERIC, unidades_vendidas_total BIGINT,
    receita_periodo_brl NUMERIC, receita_total_brl NUMERIC, data_estimada_upload TEXT,
    unidades_vendidas_total_loja BIGINT, fastmoss_produto_url TEXT,
    tiktok_shop_url TEXT, fastmoss_loja_url TEXT, status_origem TEXT,
    fonte_dados TEXT, ranking_origem INTEGER
  )
)
INSERT INTO public.products (
  produto_externo_id, nome, categoria, imagem_url, loja_nome, preco_texto,
  preco_min_brl, preco_max_brl, pais_regiao, comissao_percentual,
  unidades_vendidas_periodo, crescimento_vendas_percentual, unidades_vendidas_total,
  receita_periodo_brl, receita_total_brl, data_estimada_upload,
  unidades_vendidas_total_loja, fastmoss_produto_url, tiktok_shop_url,
  fastmoss_loja_url, status_origem, fonte_dados, ranking_origem,
  status, destaque, ordem_destaque
)
SELECT
  produto_externo_id, nome, categoria, imagem_url, loja_nome, preco_texto,
  preco_min_brl, preco_max_brl, pais_regiao, comissao_percentual,
  unidades_vendidas_periodo, crescimento_vendas_percentual, unidades_vendidas_total,
  receita_periodo_brl, receita_total_brl,
  CASE WHEN data_estimada_upload IS NULL THEN NULL
    ELSE data_estimada_upload::timestamp AT TIME ZONE 'America/Sao_Paulo' END,
  unidades_vendidas_total_loja, fastmoss_produto_url, tiktok_shop_url,
  fastmoss_loja_url, status_origem, fonte_dados, ranking_origem,
  CASE WHEN status_origem = 'Disponível' THEN 'ativo' ELSE 'pausado' END,
  false, 0
FROM catalog
ON CONFLICT (produto_externo_id) DO UPDATE SET
  nome = EXCLUDED.nome,
  categoria = EXCLUDED.categoria,
  imagem_url = EXCLUDED.imagem_url,
  loja_nome = EXCLUDED.loja_nome,
  preco_texto = EXCLUDED.preco_texto,
  preco_min_brl = EXCLUDED.preco_min_brl,
  preco_max_brl = EXCLUDED.preco_max_brl,
  pais_regiao = EXCLUDED.pais_regiao,
  comissao_percentual = EXCLUDED.comissao_percentual,
  unidades_vendidas_periodo = EXCLUDED.unidades_vendidas_periodo,
  crescimento_vendas_percentual = EXCLUDED.crescimento_vendas_percentual,
  unidades_vendidas_total = EXCLUDED.unidades_vendidas_total,
  receita_periodo_brl = EXCLUDED.receita_periodo_brl,
  receita_total_brl = EXCLUDED.receita_total_brl,
  data_estimada_upload = EXCLUDED.data_estimada_upload,
  unidades_vendidas_total_loja = EXCLUDED.unidades_vendidas_total_loja,
  fastmoss_produto_url = EXCLUDED.fastmoss_produto_url,
  tiktok_shop_url = EXCLUDED.tiktok_shop_url,
  fastmoss_loja_url = EXCLUDED.fastmoss_loja_url,
  status_origem = EXCLUDED.status_origem,
  fonte_dados = EXCLUDED.fonte_dados,
  ranking_origem = EXCLUDED.ranking_origem;
"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--migration", required=True, type=Path)
    parser.add_argument("--json", dest="json_path", type=Path)
    args = parser.parse_args()

    products, duplicates = read_rows(args.inputs)
    args.migration.parent.mkdir(parents=True, exist_ok=True)
    args.migration.write_text(
        migration_sql(products, [path.name for path in args.inputs]), encoding="utf-8"
    )
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(
            json.dumps(products, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps({"products": len(products), "duplicates_collapsed": duplicates}))


if __name__ == "__main__":
    main()
